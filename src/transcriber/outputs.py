from __future__ import annotations

import json
from dataclasses import asdict
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from openpyxl import Workbook

from .pipeline import TranscriptionResult


def _format_srt_timestamp(seconds: float) -> str:
    seconds = max(0.0, seconds)
    total_ms = int(round(seconds * 1000))
    hours = total_ms // 3_600_000
    total_ms %= 3_600_000
    minutes = total_ms // 60_000
    total_ms %= 60_000
    secs = total_ms // 1000
    millis = total_ms % 1000
    return f"{hours:02d}:{minutes:02d}:{secs:02d},{millis:03d}"


def write_txt(result: TranscriptionResult, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(result.text.strip() + "\n", encoding="utf-8")


def write_json(result: TranscriptionResult, path: Path) -> None:
    payload = asdict(result)
    payload["input_path"] = str(payload["input_path"])  # Path → str for JSON
    payload["created_at"] = datetime.now(timezone.utc).isoformat()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def write_srt(result: TranscriptionResult, path: Path) -> None:
    lines: list[str] = []
    for idx, seg in enumerate(result.segments, start=1):
        start = _format_srt_timestamp(seg["start"])
        end = _format_srt_timestamp(seg["end"])
        text = seg["text"].strip()
        lines.append(f"{idx}\n{start} --> {end}\n{text}\n")

    if not lines:
        lines.append(
            "1\n"
            "00:00:00,000 --> 00:00:05,000\n"
            f"{result.text.strip()}\n"
        )

    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")


def write_xlsx(result: TranscriptionResult, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Segments"
    ws.append(["Segment", "Start (s)", "End (s)", "Text"])

    if result.segments:
        for idx, seg in enumerate(result.segments, start=1):
            ws.append([
                idx,
                float(seg["start"]),
                float(seg["end"]),
                seg["text"].strip(),
            ])
    else:
        ws.append([1, 0.0, 0.0, result.text.strip()])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_docx(result: TranscriptionResult, path: Path) -> None:
    doc = Document()

    if result.segments:
        for seg in result.segments:
            doc.add_paragraph(str(seg["text"]).strip())
    else:
        doc.add_paragraph(result.text.strip())

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def write_outputs(
    result: TranscriptionResult,
    output_dir: Path,
    write_txt_file: bool,
    write_json_file: bool,
    write_srt_file: bool,
    write_xlsx_file: bool = False,
    write_docx_file: bool = False,
) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    stem = result.input_path.stem
    written: list[Path] = []

    if write_txt_file:
        txt_path = output_dir / f"{stem}.txt"
        write_txt(result, txt_path)
        written.append(txt_path)

    if write_json_file:
        json_path = output_dir / f"{stem}.json"
        write_json(result, json_path)
        written.append(json_path)

    if write_srt_file:
        srt_path = output_dir / f"{stem}.srt"
        write_srt(result, srt_path)
        written.append(srt_path)

    if write_xlsx_file:
        xlsx_path = output_dir / f"{stem}.xlsx"
        write_xlsx(result, xlsx_path)
        written.append(xlsx_path)

    if write_docx_file:
        docx_path = output_dir / f"{stem}.docx"
        write_docx(result, docx_path)
        written.append(docx_path)

    return written
